from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TASK_DIR = Path(__file__).resolve().parent
SOURCE_PATH = TASK_DIR / "story-draft.md"
OUTPUT_PATH = TASK_DIR / "黑化西格莉卡剧情审阅稿.docx"

ALLOWED_SPEAKERS = {"旁白", "西格莉卡", "西格莉卡？", "达妮娅", "玩家选项"}
NODE_HEADING_RE = re.compile(r"^#{3,4} `([^`]+)`$")
SPEAKER_RE = re.compile(r"^\*\*(旁白|西格莉卡？|西格莉卡|达妮娅|玩家选项)\*\*$")
CHOICE_RE = re.compile(r"^- 「(.+?)」→")


@dataclass
class SpeechBlock:
    speaker: str
    paragraphs: list[str] = field(default_factory=list)


def clean_text(text: str) -> str:
    return text.replace("`", "").strip()


def parse_nodes(markdown: str) -> dict[str, list[SpeechBlock]]:
    nodes: dict[str, list[SpeechBlock]] = {}
    current_node: str | None = None
    current_block: SpeechBlock | None = None

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        node_match = NODE_HEADING_RE.match(line)
        if node_match:
            current_node = node_match.group(1)
            nodes[current_node] = []
            current_block = None
            continue

        if current_node is None:
            continue
        if line.startswith("#"):
            current_node = None
            current_block = None
            continue
        if not line or line == "---" or line.startswith(">") or line.startswith("→"):
            continue
        if line.startswith("**演出："):
            current_block = None
            continue

        speaker_match = SPEAKER_RE.match(line)
        if speaker_match:
            speaker = speaker_match.group(1)
            if speaker == "玩家选项":
                current_block = None
            else:
                current_block = SpeechBlock(speaker)
                nodes[current_node].append(current_block)
            continue

        choice_match = CHOICE_RE.match(line)
        if choice_match:
            current_block = SpeechBlock("玩家", [clean_text(choice_match.group(1))])
            nodes[current_node].append(current_block)
            continue

        if line.startswith("**") and line.endswith("**"):
            current_block = None
            continue

        if current_block is not None:
            current_block.paragraphs.append(clean_text(line))

    return nodes


def set_run_font(run, *, ascii_font: str = "Calibri", east_asia_font: str = "Microsoft YaHei") -> None:
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia_font)


def set_style_font(style, *, ascii_font: str = "Calibri", east_asia_font: str = "Microsoft YaHei") -> None:
    style.font.name = ascii_font
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia_font)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x20, 0x21, 0x24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    speaker = doc.styles.add_style("Speaker", 1)
    set_style_font(speaker)
    speaker.font.size = Pt(10.5)
    speaker.font.bold = True
    speaker.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
    speaker.paragraph_format.space_before = Pt(8)
    speaker.paragraph_format.space_after = Pt(2)
    speaker.paragraph_format.keep_with_next = True

    dialogue = doc.styles.add_style("Dialogue", 1)
    set_style_font(dialogue)
    dialogue.font.size = Pt(11)
    dialogue.font.color.rgb = RGBColor(0x20, 0x21, 0x24)
    dialogue.paragraph_format.left_indent = Inches(0.28)
    dialogue.paragraph_format.space_before = Pt(0)
    dialogue.paragraph_format.space_after = Pt(6)
    dialogue.paragraph_format.line_spacing = 1.25
    dialogue.paragraph_format.keep_together = True

    narration = doc.styles.add_style("Narration", 1)
    set_style_font(narration)
    narration.font.size = Pt(11)
    narration.font.color.rgb = RGBColor(0x43, 0x43, 0x43)
    narration.paragraph_format.left_indent = Inches(0.28)
    narration.paragraph_format.space_before = Pt(0)
    narration.paragraph_format.space_after = Pt(6)
    narration.paragraph_format.line_spacing = 1.25
    narration.paragraph_format.keep_together = True

    choice_speaker = doc.styles.add_style("ChoiceSpeaker", 1)
    set_style_font(choice_speaker)
    choice_speaker.font.size = Pt(9.5)
    choice_speaker.font.bold = True
    choice_speaker.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
    choice_speaker.paragraph_format.space_before = Pt(0)
    choice_speaker.paragraph_format.space_after = Pt(2)
    choice_speaker.paragraph_format.keep_with_next = True

    choice_dialogue = doc.styles.add_style("ChoiceDialogue", 1)
    set_style_font(choice_dialogue)
    choice_dialogue.font.size = Pt(10.5)
    choice_dialogue.font.color.rgb = RGBColor(0x20, 0x21, 0x24)
    choice_dialogue.paragraph_format.space_before = Pt(0)
    choice_dialogue.paragraph_format.space_after = Pt(4)
    choice_dialogue.paragraph_format.line_spacing = 1.25
    choice_dialogue.paragraph_format.keep_together = True


def add_text_paragraph(container, text: str, style: str):
    paragraph = container.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_speech_block(doc: Document, block: SpeechBlock) -> None:
    if not block.paragraphs:
        return
    add_text_paragraph(doc, block.speaker, "Speaker")
    body_style = "Narration" if block.speaker == "旁白" else "Dialogue"
    for text in block.paragraphs:
        add_text_paragraph(doc, text, body_style)


def set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str = "F6F8FB") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def configure_choice_table(table) -> None:
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "D9E2F3")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for _ in range(2):
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), "4680")
        grid.append(grid_col)

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)

    for cell in table.rows[0].cells:
        cell.width = Inches(3.25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_width(cell, 4680)
        set_cell_margins(cell)
        shade_cell(cell)


def clear_cell(cell) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph._element.getparent().remove(paragraph._element)


def add_block_to_cell(cell, block: SpeechBlock) -> None:
    add_text_paragraph(cell, block.speaker, "ChoiceSpeaker")
    for text in block.paragraphs:
        add_text_paragraph(cell, text, "ChoiceDialogue")


def add_choice_table(doc: Document, columns: list[list[SpeechBlock]]) -> None:
    if len(columns) != 2:
        raise ValueError("Choice tables must contain exactly two alternatives")
    table = doc.add_table(rows=1, cols=2)
    configure_choice_table(table)
    for cell, blocks in zip(table.rows[0].cells, columns, strict=True):
        clear_cell(cell)
        for block in blocks:
            add_block_to_cell(cell, block)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)


def add_node(doc: Document, nodes: dict[str, list[SpeechBlock]], node_id: str) -> None:
    blocks = nodes.get(node_id)
    if blocks is None:
        raise KeyError(f"Missing story node: {node_id}")
    if len(blocks) == 2 and all(block.speaker == "玩家" for block in blocks):
        add_choice_table(doc, [[blocks[0]], [blocks[1]]])
        return
    for block in blocks:
        add_speech_block(doc, block)


def add_choice_with_reactions(
    doc: Document,
    nodes: dict[str, list[SpeechBlock]],
    choice_id: str,
    left_reaction_id: str,
    right_reaction_id: str,
) -> None:
    choices = nodes[choice_id]
    if len(choices) != 2:
        raise ValueError(f"Expected two choices in {choice_id}")
    add_choice_table(
        doc,
        [
            [choices[0], *nodes[left_reaction_id]],
            [choices[1], *nodes[right_reaction_id]],
        ],
    )


def add_section_gap(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_story(doc: Document, nodes: dict[str, list[SpeechBlock]]) -> None:
    stage_nodes = {
        1: ["use-1-start", "use-1-board", "use-1-borrowed-hand", "use-1-choice", "use-1-reply", "use-1-close"],
        2: ["use-2-start", "use-2-sai", "use-2-akira", "use-2-choice", "use-2-own-go", "use-2-close"],
        3: ["use-3-start", "use-3-losing", "use-3-seat", "use-3-choice", "use-3-reply", "use-3-close"],
        4: ["use-4-start", "use-4-isumi", "use-4-honda", "use-4-choice", "use-4-harsh", "use-4-next-year", "use-4-close"],
        5: ["use-5-start", "use-5-match", "use-5-hidden-move", "use-5-choice", "use-5-handoff", "use-5-close"],
        6: ["use-6-start", "use-6-loss", "use-6-stop", "use-6-isumi-return", "use-6-choice", "use-6-trace", "use-6-close"],
        7: ["use-7-start", "use-7-team", "use-7-rivals", "use-7-choice", "use-7-not-yet", "use-7-close"],
    }
    shared_effect = ["shared-effect-start", "shared-effect-eat", "shared-effect-pop", "shared-effect-unavailable"]
    for stage_number in range(1, 8):
        if stage_number > 1:
            add_section_gap(doc)
        for node_id in [*stage_nodes[stage_number], *shared_effect]:
            add_node(doc, nodes, node_id)

    add_page_break(doc)
    for node_id in [
        "corruption-start",
        "corruption-finished",
        "corruption-hikaru-losses",
        "corruption-normal-reflection",
        "corruption-normal-correction",
        "corruption-candy-choice",
        "corruption-candy-eat",
        "corruption-pop-sound",
        "corruption-dizzy",
        "corruption-question",
        "corruption-expectation-seed",
        "corruption-talent-seed",
        "corruption-strength-seed",
        "corruption-result-seed",
        "corruption-climax",
        "corruption-withdraw",
        "corruption-no-move",
        "corruption-challenge",
        "corruption-answer",
        "corruption-end",
        "duel-opening",
    ]:
        add_node(doc, nodes, node_id)

    add_page_break(doc)
    for node_id in [
        "recovery-loss-start",
        "recovery-loss-no-surprise",
        "recovery-loss-expectation",
        "recovery-loss-self-doubt",
        "recovery-loss-talent",
        "recovery-loss-unseen",
        "recovery-loss-cruelty",
        "recovery-loss-verdict",
        "recovery-loss-attempt",
        "recovery-loss-retreat",
        "recovery-loss-easier",
        "recovery-loss-fade",
        "recovery-loss-collapse",
        "recovery-loss-wake",
        "recovery-loss-wake-desk",
        "recovery-loss-wake-sugar",
        "recovery-loss-wake-denia",
        "recovery-loss-wake-choice",
        "recovery-loss-wake-reply",
        "recovery-loss-wake-end",
    ]:
        add_node(doc, nodes, node_id)

    add_page_break(doc)
    for node_id in [
        "recovery-win-start",
        "recovery-win-powerless",
        "recovery-win-unchanged",
        "recovery-win-no-miracle",
        "recovery-win-expectation-claim",
    ]:
        add_node(doc, nodes, node_id)
    add_choice_with_reactions(
        doc,
        nodes,
        "recovery-win-expectation-choice",
        "recovery-win-expectation-gentle",
        "recovery-win-expectation-firm",
    )
    add_node(doc, nodes, "recovery-win-expectation-merge")
    add_node(doc, nodes, "recovery-win-talent-claim")
    add_choice_with_reactions(
        doc,
        nodes,
        "recovery-win-talent-choice",
        "recovery-win-talent-honest",
        "recovery-win-talent-self",
    )
    add_node(doc, nodes, "recovery-win-talent-merge")
    add_node(doc, nodes, "recovery-win-strength-claim")
    add_choice_with_reactions(
        doc,
        nodes,
        "recovery-win-strength-choice",
        "recovery-win-strength-exists",
        "recovery-win-strength-witness",
    )
    add_node(doc, nodes, "recovery-win-strength-merge")
    add_node(doc, nodes, "recovery-win-cruelty-claim")
    add_choice_with_reactions(
        doc,
        nodes,
        "recovery-win-cruelty-choice",
        "recovery-win-cruelty-score",
        "recovery-win-cruelty-record",
    )
    for node_id in [
        "recovery-win-cruelty-merge",
        "recovery-win-final-question",
        "recovery-win-final-choice",
        "recovery-win-final-answer",
        "recovery-win-next-move",
        "recovery-win-release",
        "recovery-win-not-convinced",
        "recovery-win-choice-made",
        "recovery-win-fade",
        "recovery-win-collapse",
        "recovery-win-wake",
        "recovery-win-wake-desk",
        "recovery-win-wake-sugar",
        "recovery-win-wake-denia",
        "recovery-win-wake-choice",
        "recovery-win-wake-invite",
        "recovery-win-wake-reply",
        "recovery-win-wake-denia-end",
        "recovery-win-wake-end",
    ]:
        add_node(doc, nodes, node_id)


def remove_initial_empty_paragraph(doc: Document) -> None:
    if doc.paragraphs and not doc.paragraphs[0].text:
        paragraph = doc.paragraphs[0]._element
        paragraph.getparent().remove(paragraph)


def visible_text(doc: Document) -> str:
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(chunks)


def audit(doc: Document) -> None:
    text = visible_text(doc)
    forbidden = [
        "use-",
        "shared-effect",
        "corruption-",
        "recovery-",
        "审阅约束",
        "文案自检",
        "版本：",
        "尚未实装",
        "演出：",
        "→",
        "决战过程沿用",
    ]
    found = [token for token in forbidden if token in text]
    if found:
        raise ValueError(f"Non-story text leaked into DOCX: {found}")
    if not text.startswith("旁白\n午后的活动室"):
        raise ValueError("Document does not begin directly with the first narration")
    if "正式赛果仅在第八段揭晓" in text:
        raise ValueError("Review note leaked into DOCX")


def main() -> None:
    markdown = SOURCE_PATH.read_text(encoding="utf-8")
    nodes = parse_nodes(markdown)
    doc = Document()
    configure_document(doc)
    build_story(doc, nodes)
    remove_initial_empty_paragraph(doc)
    audit(doc)
    doc.core_properties.title = "黑化西格莉卡剧情审阅稿"
    doc.core_properties.subject = "剧情对白审阅"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()

from __future__ import annotations

import hashlib
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


REPLACEMENTS = {
    # 1. 相遇与借来的第一手
    "P002": "午后的阳光越过窗台，正好落在棋盘一角。西格莉卡把一本刚拆封的漫画摊在旁边，翘起的书页怎么也压不平，她索性拿棋罐抵住了它。封面上写着《棋魂》。",
    "P004": "等等，旧棋盘里真的住着一个棋士？进藤光看得见他，别人却看不见……那他平时要怎么跟佐为说话呀？",
    "P006": "佐为等了那么久，醒来以后最想做的事还是下棋。可他连一颗棋子都碰不到，得让光替他把手伸出去。",
    "T000-R0-C0-P1": "所以，是光帮他重新下棋。",
    "T000-R0-C1-P1": "光也会被他带到没见过的地方吧。",
    "P009": "嗯。佐为终于碰到了棋子，光也第一次认真看见棋盘。这样一想，好像很难说究竟是谁在帮谁呢。",
    "P011": "才刚开始就这么有意思……糟糕，我可能要一口气看下去了。先说好，等会儿不许嫌我太安静哦。",

    # 2. 被追逐的究竟是谁
    "P020": "书页已经翻过大半。西格莉卡本来读得飞快，看到塔矢亮追到学校来时，却用指腹压住纸角，又倒回去看了一遍。",
    "P022": "网上才下了几盘，大家就都在找 sai。记者在找，棋手在找，连塔矢亮也追着那股棋风跑过来了。",
    "P024": "可塔矢亮追到光面前，看到的还是佐为吧？难怪光会不服气。明明人就站在那里，对方的眼睛却越过了他。",
    "T001-R0-C0-P1": "他追的是那份强大，不一定是光。",
    "T001-R0-C1-P1": "所以光才想让他真正看自己一眼。",
    "P027": "对。光明明知道差得很远，还是要自己学、自己追。换成我，被人这样看过去……唔，我大概也咽不下这口气。",
    "P029": "总有一天，塔矢亮再回头，应该会发现追上来的那个人已经不是谁的影子了。到时候，他可不能再认错。",

    # 3. 重新坐回棋盘
    "P037": "院生排名一次次更新，西格莉卡翻页的速度也慢了。她把下巴搁在手背上，看到光又输一局，另一只手无意识地拨了拨棋罐盖。",
    "P039": "进院生以前明明已经赢了那么多，到了这里却一直输……每天都要从一群比自己强的人中间坐过去，光想想就累。",
    "P042": "可他第二天还是来了。昨天看不懂的地方再看一遍，昨天输过的人再下一次。那个位置没有因为他输棋就空着。",
    "T002-R0-C0-P1": "明知道下一盘可能还会输？",
    "T002-R0-C1-P1": "一直输的时候，坐回来才最难吧。",
    "P045": "会怕呀。最难受的还不是输，是坐下以前就知道自己可能又要输。可他还是把棋罐打开了。",
    "P047": "今天只多看懂一手，也先好好收着。明天说不定正好就用上了呢。嘿嘿，我开始替排名表紧张了。",

    # 4. 一次失误的重量
    "P058": "读到定段赛，活动室里只剩纸页摩擦的沙沙声。西格莉卡的手停在伊角认输的那一格，拇指抵着页边，很久没有翻过去。",
    "P060": "伊角明明有实力，却在最要紧的时候犯了规。认输以后，他连后面的棋也一起乱掉了……偏偏是在这一年、这几盘。",
    "P062": "本田也是。准备了那么久，最后只差一点。一年只有三个名额，门一关，外面还站着那么多人。",
    "T003-R0-C0-P1": "比赛只会记下这一次的结果。",
    "T003-R0-C1-P1": "一次失常，真的会压住一整年的努力。",
    "P065": "嗯。一年的时间那么长，结果落下来却只要一瞬间。它也不会因为谁更舍不得，就慢一点。",
    "P067": "我知道还能再考。可‘还有下一次’说起来很轻，真正要把自己从那天捡回来，哪有那么容易。",
    "P069": "如果换成我……大概会先躲到谁也找不到的地方。等呼吸顺了，再偷偷想一想，要不要回去。",

    # 5. 最强一局留下的下一手
    "P076": "佐为终于坐到塔矢名人对面时，西格莉卡一下直起身。她屏住呼吸盯着书页，连棋罐盖滚到桌边都没注意。",
    "P078": "终于下到了……他们两个找了这么久，这盘棋一定谁都不想留遗憾吧。光是看着，我手心都要出汗了。",
    "P080": "等等，最后那一手，是光看见的？佐为自己都漏过去了，他却在旁边追到了那里……",
    "T004-R0-C0-P1": "原来最强的一局也还能往下走。",
    "T004-R0-C1-P1": "佐为把光带到了自己没看见的地方。",
    "P083": "对。光只是坐在旁边看，却真的接上了他们的棋。佐为发现的时候，没有不甘心，反而那么高兴。",
    "P085": "像是一直捧在手里的东西，终于有人接住了。嗯……我很喜欢这里。再让我看一遍。",

    # 6. 消失之后继续落子
    "P093": "窗外不知什么时候安静下来。西格莉卡把漫画抱得离自己更近，翻到空下来的棋盘时，指尖在页边停了很久。",
    "P095": "佐为真的不见了。光跑了那么多地方，找棋院，找旧棋盘……什么都试过了，还是找不到。",
    "P097": "难怪他不肯再下。只要一坐到棋盘前，佐为就会从他的棋里冒出来一下。那还怎么把这一局下完呀。",
    "P099": "然后伊角回来了。他没有逼光振作，也没有说那些很厉害的话，只是坐下来，请光陪他下一局。",
    "T005-R0-C0-P1": "光终于在自己的棋里看见了佐为。",
    "T005-R0-C1-P1": "伊角只是坐下来，等他重新落子。",
    "P102": "嗯。有时候劝人往前走，反而会让人更想躲开。伊角这样陪他安安静静下一局，倒是正好。",
    "P104": "塔矢亮也终于认出来了。光的棋里有佐为留下的东西，可每一手又是光自己下的……这一次，他总算没有看错人。",

    # 7. 通往北斗杯，只到正式赛前
    "P110": "北斗杯选拔结束时，西格莉卡一下从椅背上弹起来，把漫画举到你面前。书页哗啦一响，差点碰翻棋罐。",
    "P112": "队伍终于选出来了！光和塔矢亮这次不只要互相追，还要站在同一边，去碰中国队和韩国队的年轻棋手。",
    "P114": "名单上的人都好强……有人已经出名很久，也有人第一次站到这种赛场。正式比赛还没开始，我怎么已经紧张了。",
    "T006-R0-C0-P1": "你还没看到正式比赛吧？",
    "T006-R0-C1-P1": "那最后会是哪一队赢？",
    "P117": "还没呢！我刚看到名单定下来，下一页才要进赛场。先停在这里，不许给我剧透。",
    "P119": "结局要等四周都安静的时候再看。嘿嘿，光总算走到这么大的赛场了……我想慢一点把最后这段读完。",

    # 8. 在这里才读完北斗杯
    "P123": "傍晚的活动室只开着桌边的一盏灯。达妮娅歪在远处的沙发上打盹，西格莉卡坐在棋盘前，最后一卷已经合上了，她的手却还捏着封底的一角。",
    "P125": "……看完了。日本队最后还是输了，先输给中国队，又输给韩国队。",
    "P127": "光两盘都下到了最后，也两盘都没赢。没有刚好够他翻过去的半目，也没有留给他的那一场胜利。最后一页就这样翻过去了。",
    "P129": "我一直觉得，他都走到这里了，故事总会给他一次胜利吧。结果没有。",
    "P131": "啊，我不是说前面都白费了。就是……合上书的时候，心里突然空了一块。原来拼到最后，也可能什么都接不住。",
    "P135": "谢谢。我现在正好需要一点甜的。嗯……越甜越好。",
    "P137": "她把糖含进口中。最初几声还很轻快，在过分安静的活动室里蹦了两下。",
    "P138": "啪。",
    "P139": "又一声。",
    "P140": "第三声落下时，清脆里忽然多了一点空响。噼啪声拖得越来越长，像棋子一颗颗落进没有边的房间。",
    "P142": "桌灯在西格莉卡身后晕开一圈暗红。你伸手去扶桌沿，指尖却从熟悉的木纹上滑了过去。棋盘、书页、窗外的晚霞，一层一层褪成灰色。",
    "P144": "西格莉卡？",
    "P146": "视野重新聚拢时，她还坐在原来的位置。只是那双眼睛已经染成暗红，安静得像从很久以前就在等你抬头。",
    "P147": "西格莉卡？",
    "P148": "塔矢亮第一次追上来，真的是因为光吗？……不是。他追的是借来的强大。",
    "P149": "西格莉卡？",
    "P150": "被期待听起来很温暖。可真正的你没那么强时，那双眼睛就会先变成失望。",
    "P151": "西格莉卡？",
    "P152": "伊角只失手一次，本田只差一步。名额不会问他们以前熬过多少天。",
    "P153": "西格莉卡？",
    "P154": "大家记得 sai，记得塔矢名人，记得高永夏。排在后面的人呢？连名字都不必留下。",
    "P155": "西格莉卡？",
    "P156": "北斗杯也是。两场败局压下来，前面的每一步都被盖在结果底下。",
    "P158": "你已经看见会发生什么了。还要继续吗？",
    "P160": "别看，别听，也别把手伸出去。只要还没开始，就没人能说你不够好。",
    "P162": "把手收回来吧。这样最轻松。",
    "T007-R0-C0-P1": "可我也永远不会知道结果。",
    "T007-R0-C1-P1": "如果我还是想试呢？",
    "P165": "那就试。",
    "P167": "先赢下这一局。否则，就把手收回去。",
    "P169": "她抬起手，指尖轻轻一点。黑暗从你们之间裂开，一张棋盘随之浮起。通往决战的入口亮着暗红，光落在她脸上，一动不动。",

    # 败北分支：只谈此刻的玩家与本局
    "P173": "最后一声落下后，四周忽然静了。盘面上的数字不再跳动。",
    "P174": "你输了。",
    "P175": "西格莉卡？没有笑。她把结果从头看到尾，随后抬眼看你。开口时，声音比之前更轻。",
    "P177": "你看。每一步都认真想过，还是输了。",
    "P179": "有人在等你做到。你越在意那个人，‘我没做到’就越像一句判词。",
    "P181": "其实他们还什么都没说。你已经会替他们问：是不是我还不够好？",
    "P183": "再努力一点？可你也知道，有些人走一步，就到了你要追很久的地方。",
    "P185": "视线会跟着赢的人走。你在这里坐了多久，手心出了多少汗，没有人需要知道。",
    "P187": "最后只剩这个结果。它足够让别人转身，也足够让你怀疑前面的路。",
    "P189": "所以，撑不住也没关系。你还有一条更轻松的路。",
    "P191": "你想回答，喉咙却像被什么压住。垂在身侧的手指动了一下，始终没有抬起来。",
    "P193": "停在这里吧。别再答应任何人，也别再逼自己追上谁。",
    "P195": "没有开始，就没有失败。你还可以一直相信，只是自己没有去做。",
    "P197": "她的声音慢慢退远。灰白从地面漫上来，桌椅的轮廓先化开，接着是她的脸。",
    "P199": "意识沉下去以前，你最后看见的，是自己那只始终没有抬起来的手。",
    "P201": "{username}？{username}！醒醒，能听见我吗？",
    "P203": "你猛地抬起头，额头下压着摊开的漫画。夕阳还挂在窗边。西格莉卡一手扶着桌沿，另一只手悬在你肩旁；达妮娅抱着水杯站在后面。",
    "P205": "你刚才突然就趴下去了，怎么叫都没反应。是不是低血糖？要不要去找陆医生？我陪你。",
    "P207": "我看更像昨晚又通宵了。能在西西说话的时候睡成这样，也算很有本事。",
    "P209": "我没事……只是有点累。",
    "P211": "真的？先喝点水。今天别逞强了，棋什么时候都能下。",
    "P213": "难得同意西西一次。再趴下去，可别指望我一个人把你搬到医务室。",
    "P215": "你接过水杯，杯壁的温度慢慢贴进掌心。桌边的棋盘就在视线外，你没有转头。",
    "P216": "你低头喝水。梦里那句‘把手收回来吧’，还贴在耳边。",

    # 胜利分支：四轮双选项，各自回应后立即汇合
    "P220": "最后一手落定，盘面上的光点逐个熄灭。",
    "P221": "你赢了。",
    "P222": "西格莉卡？站在对面，很久没有动。过了一会儿，她像突然忘了该怎么继续站着，慢慢坐了下来。",
    "P224": "……你赢了。",
    "P226": "可这一局外面的事没有变。明天会不会输，别人会不会失望，有没有人看见你，都还是原来的样子。",
    "P228": "这一局，替你改了什么？",
    "P230": "当所有人都说‘你一定可以’，那句话迟早会变重。做不到的时候，你最先怀疑的还是自己。",
    "T008-R0-C0-P1": "他们可以期待。我也可以说，我这次做不到。",
    "T008-R0-C0-P3": "说出口以后呢？你在意的人沉默一下，你还是会先怪自己。",
    "T008-R0-C1-P1": "我会怕他们失望。可怕，不等于我就该消失。",
    "T008-R0-C1-P3": "失望是真的。你没法装作没有看见。",
    "P233": "我会看见，也会难受。可我不替别人，把自己判成不值得。",
    "P235": "那天赋呢？有人很快就能走到你拼命追的地方。你花掉同样的时间，也未必追得上。",
    "T009-R0-C0-P1": "我追不上所有人。那就先走到我能走到的地方。",
    "T009-R0-C0-P3": "如果那个地方离你想要的还很远呢？",
    "T009-R0-C1-P1": "努力没答应让我赢。它只让我别停在原地。",
    "T009-R0-C1-P3": "只是挪动一点，也值得吗？",
    "P238": "我想自己走过以后，再回答值不值得。",
    "P240": "可人们还是会朝强者那边聚过去。没有赢的人坐在角落里，再努力也很容易被漏掉。",
    "T010-R0-C0-P1": "没人回头，我做过的事也不会倒着消失。",
    "T010-R0-C0-P3": "只剩你自己知道，你真能一直相信它？",
    "T010-R0-C1-P1": "那我先别把目光只给赢的人。",
    "T010-R0-C1-P3": "你看见一个，世界还是会漏掉更多。",
    "P243": "我管不了所有人的眼睛。至少，我不跟着一起装作没看见。",
    "P245": "结果呢？别人只要看一眼输赢，就能替你走过的那些日子下结论。",
    "T011-R0-C0-P1": "输就写输。我不删掉为了它走过的日子。",
    "T011-R0-C0-P3": "那个字会一直留着。你每次看见，都要再输一次。",
    "T011-R0-C1-P1": "这次的结果，只能说这一次。",
    "T011-R0-C1-P3": "如果下一次还是一样呢？",
    "P248": "那就等下一次来了，我再决定还走不走。",
    "P250": "你只是把痛留给以后的自己。",
    "T012-R0-C0-P1": "至少不是现在因为害怕，就替以后的我认输。",
    "T012-R0-C1-P1": "我没有赢过这个世界。我只是不想让你替我收手。",
    "P253": "即使下一次真的很难看？",
    "T013-R0-C0-P1": "难看就难看。到时候我自己看。",
    "T013-R0-C1-P1": "轮到我的时候，我自己选。",
    "P256": "……真固执。",
    "P258": "外面的事一件也没有变。我还是不知道，这样到底有什么用。",
    "P260": "可你已经把手伸出去了。我没法再说，你什么都没做。",
    "P262": "她最后一个字落下，棋盘边缘便开始失焦。暗红的天花板、桌椅和她的身影一起被白光推远，远处有人一遍遍叫你的名字。",
    "P264": "你朝声音伸手，指尖先碰到一片温热的桌面。",
    "P266": "{username}？{username}！太好了，你终于醒了！",
    "P268": "你从课桌上抬起头，眼前还浮着一层白。活动室依旧浸在夕阳里。西格莉卡半蹲在桌边，肩膀刚刚松下来；达妮娅端着水杯，在她身后打了个小小的哈欠。",
    "P270": "你刚才突然趴下去了，脸色也不太好。是不是低血糖？要不要我陪你去找陆医生？",
    "P272": "也可能是昨晚通宵。正常人不会听着糖果响两声，就当场进入休眠。",
    "P274": "我没事。只是做了个……挺累的梦。",
    "P276": "西格莉卡，陪我下一局吧。",
    "P278": "现在？你才刚醒呀。",
    "P280": "她盯着你看了两秒，又看了看桌边的棋盘。最后，她把水杯先塞进你手里，自己转身去拿棋罐。",
    "P282": "好吧。不过先喝水。下棋可以，再晕过去可不行。",
    "P284": "我就知道。水还没喝一口，棋倒先约上了……你们下吧，我去食堂占位置。",
    "P286": "第一颗棋子落在交叉点上，清脆的一声，和活动室里任何一个傍晚都一样。",
    "P287": "窗外的风翻动最后一卷漫画，达妮娅在门口催你们记得吃饭。棋盘另一边，西格莉卡正皱着鼻尖找位置。",
    "P288": "你没有催她，只把手放回膝上，等她落下下一手。",
}


COMMON_START = 303
COMMON_END = 313


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def collect_keyed_paragraphs(document: Document) -> dict[str, Paragraph]:
    keyed: dict[str, Paragraph] = {}
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            keyed[f"P{paragraph_index:03d}"] = paragraph
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    for cell_paragraph_index, paragraph in enumerate(cell.paragraphs):
                        key = (
                            f"T{table_index:03d}-R{row_index}-C{column_index}"
                            f"-P{cell_paragraph_index}"
                        )
                        keyed[key] = paragraph
            table_index += 1
    return keyed


def snapshot(document: Document) -> dict[str, object]:
    keyed = collect_keyed_paragraphs(document)
    common = [
        (
            f"P{index:03d}",
            keyed[f"P{index:03d}"].style.name,
            keyed[f"P{index:03d}"].text,
        )
        for index in range(COMMON_START, COMMON_END + 1)
    ]
    table_geometry = [
        (len(table.rows), len(table.columns), [len(cell.paragraphs) for row in table.rows for cell in row.cells])
        for table in document.tables
    ]
    styles = {key: paragraph.style.name for key, paragraph in keyed.items()}
    return {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "table_geometry": table_geometry,
        "styles": styles,
        "common": common,
    }


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    if source == output:
        raise RuntimeError("Output must not overwrite the source document")
    output.parent.mkdir(parents=True, exist_ok=True)

    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    source_document = Document(source)
    source_snapshot = snapshot(source_document)

    shutil.copyfile(source, output)
    document = Document(output)
    keyed = collect_keyed_paragraphs(document)

    missing = sorted(set(REPLACEMENTS) - set(keyed))
    if missing:
        raise RuntimeError(f"Replacement keys not found: {missing}")

    for key, text in REPLACEMENTS.items():
        if key.startswith("P") and COMMON_START <= int(key[1:]) <= COMMON_END:
            raise RuntimeError(f"Attempted to modify protected common-candy key: {key}")
        set_paragraph_text(keyed[key], text)

    # The source uses a page-break run in an otherwise empty P217. With the
    # shorter rewritten loss branch, Word places that empty paragraph alone on
    # a blank page. Move the same branch break to the end of P216 so P217 and
    # the victory heading begin together on the next page.
    for page_break in keyed["P217"]._p.xpath('.//w:br[@w:type="page"]'):
        page_break.getparent().remove(page_break)
    keyed["P218"].paragraph_format.page_break_before = False
    keyed["P216"].runs[-1].add_break(WD_BREAK.PAGE)

    document.save(output)

    result_document = Document(output)
    result_keyed = collect_keyed_paragraphs(result_document)
    result_snapshot = snapshot(result_document)

    for field in ("paragraph_count", "table_count", "table_geometry", "styles", "common"):
        if result_snapshot[field] != source_snapshot[field]:
            raise RuntimeError(f"Structure verification failed: {field}")

    for key, expected in REPLACEMENTS.items():
        if result_keyed[key].text != expected:
            raise RuntimeError(f"Text verification failed for {key}")

    if not result_keyed["P216"]._p.xpath('.//w:br[@w:type="page"]'):
        raise RuntimeError("P216 does not carry the replacement branch break")
    if result_keyed["P217"]._p.xpath('.//w:br[@w:type="page"]'):
        raise RuntimeError("P217 still contains the source page-break run")
    if result_keyed["P218"].paragraph_format.page_break_before is True:
        raise RuntimeError("P218 still forces an additional page break")

    if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("Source document changed during rewrite")

    print(f"source_sha256={source_hash}")
    print(f"output_sha256={hashlib.sha256(output.read_bytes()).hexdigest()}")
    print(f"replacements={len(REPLACEMENTS)}")
    print("structure_and_styles=unchanged")
    print("common_candy_text_and_styles=unchanged")
    print(output)


if __name__ == "__main__":
    main()

from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_body_blocks(document: Document):
    paragraph_index = 0
    table_index = 0
    for body_index, child in enumerate(document.element.body.iterchildren()):
        if child.tag == qn("w:p"):
            yield {
                "kind": "paragraph",
                "body_index": body_index,
                "paragraph_index": paragraph_index,
                "paragraph": Paragraph(child, document),
            }
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            yield {
                "kind": "table",
                "body_index": body_index,
                "table_index": table_index,
                "table": Table(child, document),
            }
            table_index += 1


def main() -> None:
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    document = Document(source)
    records: list[dict[str, object]] = []

    for block in iter_body_blocks(document):
        if block["kind"] == "paragraph":
            paragraph = block.pop("paragraph")
            records.append(
                {
                    **block,
                    "key": f"P{block['paragraph_index']:03d}",
                    "style": paragraph.style.name if paragraph.style else "",
                    "text": paragraph.text,
                }
            )
            continue

        table = block.pop("table")
        table_record: dict[str, object] = {
            **block,
            "key": f"T{block['table_index']:03d}",
            "rows": len(table.rows),
            "columns": len(table.columns),
            "cells": [],
        }
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                for cell_paragraph_index, paragraph in enumerate(cell.paragraphs):
                    table_record["cells"].append(
                        {
                            "key": (
                                f"T{block['table_index']:03d}"
                                f"-R{row_index}-C{column_index}-P{cell_paragraph_index}"
                            ),
                            "row": row_index,
                            "column": column_index,
                            "paragraph_index": cell_paragraph_index,
                            "style": paragraph.style.name if paragraph.style else "",
                            "text": paragraph.text,
                        }
                    )
        records.append(table_record)

    payload = {
        "source": str(source),
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "records": records,
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()

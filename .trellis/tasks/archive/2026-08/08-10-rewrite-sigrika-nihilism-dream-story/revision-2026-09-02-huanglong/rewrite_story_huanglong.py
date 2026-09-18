from __future__ import annotations

import hashlib
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


REPLACEMENTS = {
    # 1. 旧棋盘里的客人
    "P002": "午后的活动室被晒得暖烘烘的。西格莉卡拿棋罐压住一本新拆封的小说，自己趴在桌沿看得入神。封面画着一方嵌有黑石的旧棋盘，书名是《星枰旧响》。",
    "P004": "诶，等一下。江照只是擦了擦棋盘，里面的残响就醒了？这个岑观澜，隔了那么多年，开口第一句居然是在嫌他把星子摆错……",
    "P006": "可岑观澜自己碰不到棋子，只能让江照替他落。江照本来还想把棋盘卖掉修屋顶呢，被他念得只好坐下了。唔，这到底算谁帮谁呀？",
    "T000-R0-C0-P1": "江照借给他一双手。",
    "T000-R0-C1-P1": "岑观澜也替江照打开了一条路。",
    "P009": "对，他们都占到便宜了嘛。岑观澜借到一双手，江照呢——原本满脑子只有漏雨的屋顶，现在居然会为了下一手该落在哪儿跟人生气。",
    "P011": "这才刚见面，他们已经吵了三次了……嘿嘿，我赌下一章还要吵。先说好，要是我看得忘了时间，你记得叫我一声。",

    # 2. 被目光越过的人
    "P020": "书页已经翻过大半。西格莉卡原本看得飞快，读到祁珩连夜赶来今州时，却按住页角，倒回去又看了一遍。",
    "P022": "那盘棋才被传到终端上，明庭那边就全在找下棋的人。祁珩也是，连第二天的公开课都推了，赶最早一班城际车来了今州。",
    "P024": "可他见到江照以后，问的每一句都是那盘棋。等他发现江照根本没有那么强，脸上的失望……唔。江照明明就站在他面前，他却一直在看另一个人。",
    "T001-R0-C0-P1": "他追来今州，是为了那份强大。",
    "T001-R0-C1-P1": "难怪江照想让他重新看自己。",
    "P027": "换成我，我大概也会憋着一口气。不是非要让祁珩道歉，就是……至少有一天，他叫出名字的时候，别再叫错人。",
    "P029": "江照真的去报名棋馆了。明知道要从头学，会走得很慢，他还是去了。好，那我也先不替他着急——再看一章，就一章。",

    # 3. 同一把椅子
    "P037": "窗外下起了雨，水珠沿玻璃拖出一道道细线。书里的江照刚进今州棋馆，一天连输六盘；书外的西格莉卡也不知不觉坐直了。",
    "P039": "第六盘结束以后，连陪练都不知道该怎么安慰他。他走的时候还装得挺轻松，一出门就把伞撞反了……看得我都想假装没看见。",
    "P042": "可第二天，那把椅背上又搭着他的湿外套。他没说昨天的事，只把棋谱摊开，问第一盘是从哪一手开始坏掉的。",
    "T002-R0-C0-P1": "他不怕今天还是输吗？",
    "T002-R0-C1-P1": "一直输的时候，回来才最难。",
    "P045": "当然怕呀。要是我，走到门口可能会先绕两圈，再装作只是路过……可他还是进去了。",
    "P047": "而且到了第四盘，他终于让对面多想了一会儿。只有一小会儿也算嘛。来，下一页，让我看看他今天能不能少输一盘。",

    # 4. 三枚执印
    "P058": "读到取印试，活动室里只剩雨声。西格莉卡的拇指抵在薛宁伸手去挪那颗星子的画面上，迟迟没有翻页。",
    "P060": "薛宁明明知道，落定的棋不能拿回去。可那一瞬间，他的手像没听见脑子说话……裁定一下来，后面几盘也全乱了。",
    "P062": "乔野更难受。他没有犯错，最后只排第四。棋馆一年就发三枚执印，名单薄薄一张，偏偏装不下第四个人。",
    "T003-R0-C0-P1": "裁定只看这一刻发生了什么。",
    "T003-R0-C1-P1": "准备了那么久，也可能被一次失常压住。",
    "P065": "嗯。那些天明明一天天熬得那么慢，名单贴出来却只要一下。纸角被风吹起来，名字还是不会往下多出一行。",
    "P067": "书里的人都说还有明年。可‘明年’两个字，说出口一点也不重，真要背着它走完下一年，就不是一回事了。",
    "P069": "而且大家越相信你能做到，门关上的时候就越不敢回头。你还得先想，该怎么让等你的人别那么失望……唉，薛宁现在一定谁都不想见。",

    # 5. 被接住的一手
    "P076": "岑观澜终于等到与裴鹤生对局的那天。盘古终端的微光铺满插图，窗外是今州入夜后的灯火。西格莉卡屏住呼吸，连滚到手边的棋罐盖也没顾上。",
    "P078": "他们隔着这么久才碰上，谁都不肯先退。岑观澜刚才还一直催江照快一点，现在反而每一手都要想好久……",
    "P080": "等一下。最后那条变化，是江照先看见的？岑观澜和裴鹤生都以为已经结束了，他却还顺着盘面往后追了一手。",
    "T004-R0-C0-P1": "原来这样的棋也还没走到尽头。",
    "T004-R0-C1-P1": "岑观澜把他带到了自己看漏的地方。",
    "P083": "岑观澜发现以后，笑得比赢棋还开心。江照反倒被他笑得发毛，一直问自己是不是算错了……真是的，这次可不是。",
    "P085": "他等了那么久，最后等到的也许不只是一盘棋。有人能从他停下的地方再往前看，这样他就不用把所有东西一起带走了。",

    # 6. 没有回答的棋盘
    "P093": "雨不知什么时候停了。西格莉卡把小说抱得离自己更近，读到旧棋盘彻底安静下来时，指尖也跟着停在了页边。",
    "P095": "岑观澜真的不见了。江照把棋盘送去华胥研究院，守着检测仪等到半夜，终端上还是只有一条冷冰冰的‘未发现有效频率’。",
    "P097": "难怪他不肯再下。每一颗星子落下去，都像在等那个人嫌他太急、太慢、又看漏了什么……可这次没有回答。",
    "P099": "后来薛宁从重州回来了。他没劝江照，也没问那张空了好久的棋桌，只把白子倒进棋罐，说自己在外面憋了一路，缺个对手。",
    "T005-R0-C0-P1": "江照终于又把第一颗棋子拿起来了。",
    "T005-R0-C1-P1": "薛宁只是坐在那里，等他愿意开始。",
    "P102": "这样刚刚好。要是有人围着他说‘振作一点’，他肯定又要躲。薛宁只管等，等到那颗棋子在手里捂热了，再陪他下一盘。",
    "P104": "祁珩也认出来了。那盘棋里有岑观澜留下的习惯，可那个会在角上犹豫、又突然抢先的人，只能是江照。这一次，他叫的是江照的名字。",

    # 7. 开赛以前
    "P110": "七宿青年联弈的选拔名单公布时，西格莉卡一下从椅背上弹起来，把小说举到你面前。书页哗啦一响，差点碰翻桌上的棋罐。",
    "P112": "进了！江照和祁珩都进了今州队。决赛阶段要碰明庭和重州，他们这回不只互相追，还得坐在同一边。",
    "P114": "明庭那位主将从小就出名，重州队也几乎全是执印棋手。江照第一次去这么大的赛场，出发前还把通行凭证落在棋馆里……我怎么比他还紧张。",
    "T006-R0-C0-P1": "你还没看到正式比赛吧？",
    "T006-R0-C1-P1": "那最后是哪一队赢了？",
    "P117": "还没有！现在才刚到开幕前夜，下一章才进赛场。别往后翻，也不许看目录，我要自己看。",
    "P119": "最后这几章得留到安静的时候。好不容易走到这里，不能三两下看完……我明天把结局带来。",

    # 8. 结局，以及梦境入口
    "P123": "第二天傍晚，活动室只开了桌边一盏灯。达妮娅歪在远处的沙发上打盹，西格莉卡坐在棋盘前，最后一页已经翻完，她的手却还捏着封底的一角。",
    "P125": "……看完了。今州第一轮输给明庭，第二轮又输给重州。三方正赛，两轮都没有赢。",
    "P127": "江照也输了两盘。第一盘追到最后还是差一点，第二盘从中段起就再没追上。书里没有临时加赛，也没有谁把一场胜利让给他。",
    "P129": "我一直以为，都写到最后一卷了，总会留一场给他吧。哪怕只赢一盘也好……结果真的没有。",
    "P131": "书还是很好看。我只是有点……不知道该把这口气放到哪里。原来走到故事最后，也不一定有一场胜利在那儿等着他。",
    "P135": "谢谢。现在确实需要一点甜的……就一颗。达妮娅睡着了，你别告诉她我又偷吃。",
    "P137": "她把糖含进口中。起初几声还很轻快，在安静的活动室里跳了两下。",
    "P138": "啪。",
    "P139": "又一声。",
    "P140": "第三声拖得很长。细碎的噼啪渐渐沉下去，像一颗颗棋子落进空屋，落下以后，许久也等不到回音。",
    "P142": "桌灯在西格莉卡身后洇开一圈暗红。你下意识去扶桌沿，指腹却像隔着一层水，怎么也碰不到熟悉的木纹。窗、棋盘、她手边合起的书，一样样褪了颜色。",
    "P144": "西格莉卡？",
    "P146": "视野重新聚拢时，她仍坐在原来的位置。那双眼睛却染成暗红，脸上没有方才读完结局的错愕，只有一种近乎疲倦的平静。",
    "P147": "西格莉卡？",
    "P148": "祁珩为什么赶去今州？因为江照吗？……不是。只因为他借来的那盘棋足够强。",
    "P149": "西格莉卡？",
    "P150": "后来所有人都叫他‘今州的希望’。输掉以后，他做的第一件事却是道歉。期待就是这样压上去的。",
    "P151": "西格莉卡？",
    "P152": "薛宁的手只错了一次，乔野也只少一个名次。名单不会问，他们在那扇门外站了多少年。",
    "P153": "西格莉卡？",
    "P154": "人群总会先挤到强者身边。终端首页留下的是名手棋谱。至于被淘汰的人，连那天坐在哪张桌前都没人关心。",
    "P155": "西格莉卡？",
    "P156": "江照已经走得够远了，还是连输两盘。差距没有因为他是故事里的主角，就替他让路。",
    "P158": "你都看见了。努力没有接住他，期待没有，想被人看见也没有。还要继续吗？",
    "P160": "不去看，不去听，也别把手伸出去。只要你没有交出答案，就没人能拿着结果回来，告诉你不够好。",
    "P162": "把手收好吧。会轻松一点。",
    "T007-R0-C0-P1": "可那样，我也永远不会知道。",
    "T007-R0-C1-P1": "如果我还是想把手伸出去呢？",
    "P165": "那就来。",
    "P167": "先让我看看，你能不能守住刚才那句话。否则，就停在这里。",
    "P169": "她抬起手，指尖在空中轻轻一点。黑暗从你们之间裂开，棋盘随之浮起。暗红的光掠过她的脸，她没有催促，只是等你坐下。",

    # 败北分支：只谈玩家与此刻，不再引用小说
    "P173": "最后一声落下，盘面上的光点逐个熄灭。四周静得只剩你的呼吸。",
    "P174": "你输了。",
    "P175": "西格莉卡？没有笑。她低头看了一遍结果，像是在确认什么，随后把手从棋盘边缘收了回去。",
    "P177": "刚才每一步，你都想过。结果没有因此对你客气一点。",
    "P179": "外面还有人在等你把事情做好。现在回去，最难受的甚至不是责怪——是他们还会安慰你，说没关系。",
    "P181": "你会盯着那句话想：他们是不是早就知道，我其实做不到。",
    "P183": "你当然还能更努力。可有些人走一步，就到了你要追很久的地方。追不上，不是再熬一晚就会消失的事。",
    "P185": "目光也会跟着赢的人走。你在这里坐了多久，手心出了多少汗，不会有人特意回来问。",
    "P187": "最后留下来的只有这个结果。一个字，就够你把前面的日子重新怀疑一遍。",
    "P189": "所以，到这里就够了。你只是累了。停下来，不用再逼自己。",
    "P191": "你想开口，喉咙却像压着一团浸了水的棉花。垂在身侧的手指蜷了一下，始终没有抬起来。",
    "P193": "别再答应谁，也不用证明什么。没有人看见，就没有人能失望。",
    "P195": "只要还没开始，你就永远可以相信——也许你本来做得到，只是没有去做。",
    "P197": "她的声音慢慢退远。灰白从地面漫上来，先吞掉桌脚，再吞掉她的肩膀。",
    "P199": "意识沉下去以前，你最后看见的，是自己那只始终没有伸出去的手。",
    "P201": "{username}？{username}！醒醒，能听见我吗？",
    "P203": "你猛地抬起头，手臂被额头压得发麻。夕阳还停在窗边。西格莉卡扶着桌沿，另一只手悬在你肩旁；达妮娅端着水杯站在后面。",
    "P205": "你刚才突然就趴下去了，怎么叫都没反应。是不是低血糖？要不要去找陆医生？我陪你。",
    "P207": "我看更像昨晚又没睡。能在西西说话的时候睡成这样，也算一种本事。",
    "P209": "我没事……就是有点累。",
    "P211": "真的？先喝水。今天别逞强了，剩下的事明天再说。",
    "P213": "难得同意西西一次。再趴下去，可别指望我一个人把你搬到医务室。",
    "P215": "你接过水杯，温度一点点贴进掌心。桌边的棋盘就在余光里，你把视线移向窗外。",
    "P216": "你低头喝水。梦里那句‘把手收好吧’，还贴在耳边，没有完全散去。",

    # 胜利分支：四轮选择立即汇合
    "P220": "最后一手落定，盘面上的暗红一格格熄灭。",
    "P221": "你赢了。",
    "P222": "西格莉卡？站在对面，很久没有动。过了一会儿，她扶住椅背。椅脚在地面上蹭出一声，她坐下时，目光还停在盘面上。",
    "P224": "……你赢了。",
    "P226": "可出去以后，那些事一件也不会少。你还是可能做不到，可能被忽略，也可能让等你的人失望。",
    "P228": "这一局，替你改了什么？",
    "P230": "当别人说‘你一定可以’，那句话迟早会变重。真做不到的时候，你第一个怪的还是自己。",
    "T008-R0-C0-P1": "他们可以失望。我也可以承认，这次我做不到。",
    "T008-R0-C0-P3": "真说出口呢？对方只要沉默一下，你就会想把话收回去。",
    "T008-R0-C1-P1": "我会怕。可怕他们失望，不等于我就该消失。",
    "T008-R0-C1-P3": "失望是真的。你没办法假装看不见。",
    "P233": "我会看见，也许还会躲一会儿。可我不替任何人，把失望写成‘我不值得’。",
    "P235": "那天赋呢？有人走一步，就到了你要追很久的地方。你花掉同样的时间，也未必追得上。",
    "T009-R0-C0-P1": "追不上就是追不上。可我想先看看，自己能走到哪里。",
    "T009-R0-C0-P3": "如果走到那里，还是离你想要的很远呢？",
    "T009-R0-C1-P1": "努力没答应让我赢。我也没答应现在就停。",
    "T009-R0-C1-P3": "只是往前挪一点，也值得花掉那些时间？",
    "P238": "我都还没走到那里，怎么会知道？总不能现在怕一怕，就算替以后的自己去过了。",
    "P240": "可人们还是会朝强者那边挤。没赢的人坐在角落里，再认真也很容易被漏掉。",
    "T010-R0-C0-P1": "没人看见，也不会把我花过的时间退回来。",
    "T010-R0-C0-P3": "只剩你自己记得。你能一直靠自己记着吗？",
    "T010-R0-C1-P1": "那我先别把自己的眼睛只给赢的人。",
    "T010-R0-C1-P3": "你看见一个，世界还是会漏掉更多。",
    "P243": "我管不了所有人看谁。至少，我不跟着一起把没被看见的部分删掉。",
    "P245": "结果呢？别人只要看一眼输赢，就能替你走过的那些日子下结论。",
    "T011-R0-C0-P1": "输就写输。别顺手把我走到这里的日子也划掉。",
    "T011-R0-C0-P3": "可那个字会留下。你每次看见，都得再输一次。",
    "T011-R0-C1-P1": "这个结果只管这一次，管不了我以后怎么选。",
    "T011-R0-C1-P3": "如果下一次还是一样呢？",
    "P248": "等下一次真的来了，我再决定。现在我不知道，也不想抢着替那时的自己认输。",
    "P250": "你只是把疼留给了以后的自己。",
    "T012-R0-C0-P1": "也许吧。可现在把手收回去，那些害怕也不会凭空不见。",
    "T012-R0-C1-P1": "我没赢过这些事。我只是不让你替我收手。",
    "P253": "即使下一次更难看？",
    "T013-R0-C0-P1": "难看就难看。到时候，我自己看。",
    "T013-R0-C1-P1": "轮到我的时候，我自己选。",
    "P256": "……你还是把这只手，从我这里拿回去了。",
    "P258": "我以为，只要劝你别伸手，至少这一回，你就不会再疼。",
    "P260": "可你非要自己选。我拦不住。",
    "P262": "她最后一个字落下，棋盘边缘便开始失焦。暗红的天花板、桌椅和她的身影一起被白光推远。很远的地方，有人一遍遍叫你的名字。",
    "P264": "你朝声音伸出手，指尖先碰到一片温热的桌面。",
    "P266": "{username}？{username}！太好了，你终于醒了！",
    "P268": "你从课桌上抬起头，眼前还浮着一层白。活动室仍泡在夕阳里。西格莉卡半蹲在桌边，绷紧的肩膀刚刚松下来；达妮娅端着水杯，在她身后打了个小小的哈欠。",
    "P270": "你刚才突然趴下去了，脸色也不太好。是不是低血糖？要不要我陪你去找陆医生？",
    "P272": "也可能是昨晚通宵。正常人不会听着糖果响两声，就当场睡得叫不醒。",
    "P274": "我没事。只是做了个……挺累的梦。",
    "P276": "西格莉卡，陪我下一局吧。",
    "P278": "现在？你才刚醒呀……",
    "P280": "她盯着你看了两秒，又看向桌边的棋盘。最后，她先把水杯塞进你手里，自己转身去拿棋罐。",
    "P282": "好吧。不过先喝水。下棋可以，再晕过去可不行。",
    "P284": "我就知道。水还没喝一口，棋倒先约上了……你们下吧，我去食堂占位置。",
    "P286": "第一颗棋子落在交叉点上，清脆的一声，和活动室里任何一个傍晚都一样。",
    "P287": "窗帘被风拱起，又慢慢落下。达妮娅已经走到门口，还在催你们记得吃饭。",
    "P288": "棋盘另一边，西格莉卡皱着鼻尖找了半天位置。你没有催，只把手放回膝上，等她落下下一手。",
}


COMMON_START = 303
COMMON_END = 313

BANNED_ANYWHERE = (
    "棋魂",
    "进藤光",
    "佐为",
    "塔矢",
    "伊角",
    "本田",
    "北斗杯",
    "sai",
    "Sai",
    "前辈",
    "项目",
    "发布",
    "观众",
    "脚本",
    "NPC",
    "第四墙",
)

POST_BATTLE_BANNED = (
    "江照",
    "岑观澜",
    "祁珩",
    "薛宁",
    "乔野",
    "裴鹤生",
    "星枰",
    "七宿青年联弈",
    "取印试",
    "执印",
    "今州",
    "明庭",
    "重州",
    "华胥研究院",
    "盘古终端",
)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def collect_keyed_paragraphs(document: Document) -> dict[str, Paragraph]:
    keyed: dict[str, Paragraph] = {}
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            keyed[f"P{paragraph_index:03d}"] = paragraph
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    for cell_paragraph_index, paragraph in enumerate(cell.paragraphs):
                        key = (
                            f"T{table_index:03d}-R{row_index}-C{column_index}"
                            f"-P{cell_paragraph_index}"
                        )
                        keyed[key] = paragraph
            table_index += 1
    return keyed


def snapshot(document: Document) -> dict[str, object]:
    keyed = collect_keyed_paragraphs(document)
    common = [
        (
            f"P{index:03d}",
            keyed[f"P{index:03d}"].style.name,
            keyed[f"P{index:03d}"].text,
        )
        for index in range(COMMON_START, COMMON_END + 1)
    ]
    table_geometry = [
        (
            len(table.rows),
            len(table.columns),
            [len(cell.paragraphs) for row in table.rows for cell in row.cells],
        )
        for table in document.tables
    ]
    styles = {key: paragraph.style.name for key, paragraph in keyed.items()}
    return {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "table_geometry": table_geometry,
        "styles": styles,
        "common": common,
    }


def visible_text(keyed: dict[str, Paragraph]) -> str:
    return "\n".join(paragraph.text for paragraph in keyed.values())


def assert_content_contract(keyed: dict[str, Paragraph]) -> None:
    all_text = visible_text(keyed)
    for term in BANNED_ANYWHERE:
        if term in all_text:
            raise RuntimeError(f"Banned term remains in document: {term}")

    section_7 = "\n".join(keyed[f"P{index:03d}"].text for index in range(110, 120))
    if any(term in section_7 for term in ("输了", "败给", "两轮都没有赢", "两战皆负")):
        raise RuntimeError("Section 7 reveals the final result")

    section_8 = "\n".join(keyed[f"P{index:03d}"].text for index in range(123, 170))
    required_section_8 = ("今州第一轮输给明庭", "第二轮又输给重州", "江照也输了两盘")
    for term in required_section_8:
        if term not in section_8:
            raise RuntimeError(f"Section 8 result is incomplete: {term}")

    post_battle = "\n".join(
        keyed[f"P{index:03d}"].text for index in range(173, 289)
    )
    post_battle += "\n" + "\n".join(
        paragraph.text
        for key, paragraph in keyed.items()
        if key.startswith("T") and int(key[1:4]) >= 8
    )
    for term in POST_BATTLE_BANNED:
        if term in post_battle:
            raise RuntimeError(f"Novel-specific term remains after battle: {term}")

    for table_index in range(8, 12):
        for column_index in range(2):
            for paragraph_index in range(4):
                key = (
                    f"T{table_index:03d}-R0-C{column_index}-P{paragraph_index}"
                )
                if key not in keyed or not keyed[key].text.strip():
                    raise RuntimeError(f"Victory response cell is incomplete: {key}")

    for convergence_key in ("P233", "P238", "P243", "P248"):
        if not keyed[convergence_key].text.strip():
            raise RuntimeError(f"Victory convergence is empty: {convergence_key}")


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    if source == output:
        raise RuntimeError("Output must not overwrite the source document")
    output.parent.mkdir(parents=True, exist_ok=True)

    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    source_document = Document(source)
    source_snapshot = snapshot(source_document)

    shutil.copyfile(source, output)
    document = Document(output)
    keyed = collect_keyed_paragraphs(document)

    missing = sorted(set(REPLACEMENTS) - set(keyed))
    if missing:
        raise RuntimeError(f"Replacement keys not found: {missing}")

    for key, text in REPLACEMENTS.items():
        if key.startswith("P") and COMMON_START <= int(key[1:]) <= COMMON_END:
            raise RuntimeError(f"Attempted to modify protected common-candy key: {key}")
        set_paragraph_text(keyed[key], text)

    # Keep the branch break, but attach it to the final loss paragraph so Word
    # cannot strand the source's otherwise-empty break paragraph on a page.
    for page_break in keyed["P217"]._p.xpath('.//w:br[@w:type="page"]'):
        page_break.getparent().remove(page_break)
    keyed["P218"].paragraph_format.page_break_before = False
    keyed["P216"].runs[-1].add_break(WD_BREAK.PAGE)

    document.save(output)

    result_document = Document(output)
    result_keyed = collect_keyed_paragraphs(result_document)
    result_snapshot = snapshot(result_document)

    for field in ("paragraph_count", "table_count", "table_geometry", "styles", "common"):
        if result_snapshot[field] != source_snapshot[field]:
            raise RuntimeError(f"Structure verification failed: {field}")

    for key, expected in REPLACEMENTS.items():
        if result_keyed[key].text != expected:
            raise RuntimeError(f"Text verification failed for {key}")

    if not result_keyed["P216"]._p.xpath('.//w:br[@w:type="page"]'):
        raise RuntimeError("P216 does not carry the replacement branch break")
    if result_keyed["P217"]._p.xpath('.//w:br[@w:type="page"]'):
        raise RuntimeError("P217 still contains the source page-break run")
    if result_keyed["P218"].paragraph_format.page_break_before is True:
        raise RuntimeError("P218 still forces an additional page break")

    assert_content_contract(result_keyed)

    if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("Source document changed during rewrite")

    print(f"source_sha256={source_hash}")
    print(f"output_sha256={hashlib.sha256(output.read_bytes()).hexdigest()}")
    print(f"replacements={len(REPLACEMENTS)}")
    print("structure_and_styles=unchanged")
    print("common_candy_text_and_styles=unchanged")
    print("content_contract=passed")
    print(output)


if __name__ == "__main__":
    main()
